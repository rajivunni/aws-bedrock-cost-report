#!/usr/bin/env python3
"""Monthly selected-service costs in one authenticated AWS billing context."""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path


class DataError(ValueError):
    """An incomplete or inconsistent source response, without raw payloads."""


def shift_month(month: date, offset: int) -> date:
    ordinal = month.year * 12 + month.month - 1 + offset
    year, index = divmod(ordinal, 12)
    return date(year, index + 1, 1)


def reporting_period(months: int, end_month: str | None = None,
                     today: date | None = None) -> tuple[str, str]:
    """Return an inclusive start and exclusive end for complete calendar months."""
    if not 1 <= months <= 13:
        raise ValueError("months must be between 1 and 13")
    current = (today or date.today()).replace(day=1)
    end = current
    if end_month is not None:
        if not re.fullmatch(r"\d{4}-\d{2}", end_month):
            raise ValueError("end-month must use YYYY-MM")
        end = shift_month(date.fromisoformat(end_month + "-01"), 1)
        if end > current:
            raise ValueError("end-month must be a completed month")
    return shift_month(end, -months).isoformat(), end.isoformat()


def expected_months(start: str, end: str) -> list[str]:
    first, stop = date.fromisoformat(start), date.fromisoformat(end)
    if first.day != 1 or stop.day != 1 or first >= stop:
        raise ValueError("report period must use increasing first-of-month dates")
    result = []
    while first < stop:
        result.append(first.isoformat())
        first = shift_month(first, 1)
    return result


def validate_accounts(values) -> list[str]:
    accounts = []
    for value in values:
        if not isinstance(value, str) or not re.fullmatch(r"[0-9]{12}", value.strip()):
            raise ValueError("each linked account ID must be text containing exactly 12 digits")
        account = value.strip()
        if account in accounts:
            raise ValueError("duplicate linked account ID")
        accounts.append(account)
    if not accounts:
        raise ValueError("at least one linked account ID is required")
    return accounts


def load_accounts(path: Path) -> list[str]:
    """A header prevents accidentally treating payer IDs or unrelated columns as input."""
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as source:
            reader = csv.DictReader(source)
            if reader.fieldnames != ["linked_account_id"]:
                raise ValueError("CSV must have exactly one column named linked_account_id")
            rows = list(reader)
            if any(None in row for row in rows):
                raise ValueError("CSV rows must contain exactly one value")
            return validate_accounts(row["linked_account_id"] for row in rows)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            rows = list(book.active.iter_rows(values_only=True))
            if not rows or rows[0] != ("linked_account_id",):
                raise ValueError("XLSX active sheet must have one column named linked_account_id")
            return validate_accounts(row[0] for row in rows[1:] if any(v is not None for v in row))
        finally:
            book.close()
    raise ValueError("accounts file must be .csv or .xlsx")


def pages(method, **request):
    """Both CE methods use NextPageToken. Reject cycles and unbounded pagination."""
    seen = set()
    for _ in range(1000):
        response = method(**request)
        if not isinstance(response, dict):
            raise DataError("invalid_response")
        yield response
        token = response.get("NextPageToken")
        if not token:
            return
        if not isinstance(token, str) or token in seen:
            raise DataError("invalid_pagination")
        seen.add(token)
        request["NextPageToken"] = token
    raise DataError("pagination_limit")


def discover_services(ce, account: str, keywords: list[str], start: str, end: str) -> list[str]:
    matches = set()
    for response in pages(
        ce.get_dimension_values, TimePeriod={"Start": start, "End": end},
        Dimension="SERVICE", Context="COST_AND_USAGE",
        Filter={"Dimensions": {"Key": "LINKED_ACCOUNT", "Values": [account]}},
    ):
        values = response.get("DimensionValues")
        if not isinstance(values, list):
            raise DataError("missing_service_dimensions")
        for item in values:
            value = item.get("Value") if isinstance(item, dict) else None
            if not isinstance(value, str) or not value:
                raise DataError("invalid_service_dimension")
            if any(keyword.casefold() in value.casefold() for keyword in keywords):
                matches.add(value)
    return sorted(matches)


@dataclass(frozen=True)
class CostRow:
    account: str
    month: str
    service: str
    amount: Decimal
    unit: str
    estimated: bool


@dataclass(frozen=True)
class Coverage:
    account: str
    month: str
    status: str
    detail: str = ""


@dataclass
class Report:
    start: str
    end: str
    accounts: list[str]
    keywords: list[str]
    rows: list[CostRow]
    coverage: list[Coverage]
    synthetic: bool = False

    @property
    def complete(self) -> bool:
        return bool(self.coverage) and all(row.status == "COMPLETE" for row in self.coverage)


def query_costs(ce, account: str, services: list[str], start: str, end: str) -> tuple[list[CostRow], set[str]]:
    result, observed, unique = [], set(), set()
    expected = set(expected_months(start, end))
    for response in pages(
        ce.get_cost_and_usage, TimePeriod={"Start": start, "End": end},
        Granularity="MONTHLY", Metrics=["UnblendedCost"],
        Filter={"And": [
            {"Dimensions": {"Key": "LINKED_ACCOUNT", "Values": [account]}},
            {"Dimensions": {"Key": "SERVICE", "Values": services}},
        ]},
        GroupBy=[{"Type": "DIMENSION", "Key": "SERVICE"}],
    ):
        periods = response.get("ResultsByTime")
        if not isinstance(periods, list):
            raise DataError("missing_cost_periods")
        for period in periods:
            try:
                month = period["TimePeriod"]["Start"]
                stop = period["TimePeriod"]["End"]
                groups, estimated = period["Groups"], period["Estimated"]
                if (month not in expected or stop != shift_month(date.fromisoformat(month), 1).isoformat()
                        or not isinstance(groups, list) or not isinstance(estimated, bool)):
                    raise DataError("invalid_cost_period")
                observed.add(month)
                for group in groups:
                    keys, metric = group["Keys"], group["Metrics"]["UnblendedCost"]
                    if not isinstance(keys, list) or len(keys) != 1 or keys[0] not in services:
                        raise DataError("invalid_service_group")
                    amount, unit = Decimal(metric["Amount"]), metric["Unit"]
                    if not amount.is_finite() or not isinstance(unit, str) or not unit:
                        raise DataError("invalid_cost_metric")
                    key = (month, keys[0], unit)
                    if key in unique:
                        raise DataError("duplicate_cost_group")
                    unique.add(key)
                    result.append(CostRow(account, month, keys[0], amount, unit, estimated))
            except (KeyError, TypeError, ValueError, InvalidOperation) as error:
                if isinstance(error, DataError):
                    raise
                raise DataError("invalid_cost_response") from error
    return result, observed


def collect_report(ce, accounts: list[str], keywords: list[str], start: str, end: str,
                   synthetic: bool = False) -> Report:
    accounts = validate_accounts(accounts)
    if not keywords or any(not isinstance(k, str) or not k.strip() for k in keywords):
        raise ValueError("at least one nonempty service keyword is required")
    months = expected_months(start, end)
    report = Report(start, end, accounts, keywords, [], [], synthetic)
    for account in accounts:
        try:
            services = discover_services(ce, account, keywords, start, end)
            if not services:
                report.coverage.extend(Coverage(account, m, "NO_MATCHING_SERVICES") for m in months)
                continue
            rows, observed = query_costs(ce, account, services, start, end)
            # Keep an account only after all pages validate. A failed later page cannot understate its total.
            report.rows.extend(rows)
            populated = {row.month for row in rows}
            for month in months:
                status = "COMPLETE" if month in populated else (
                    "NO_COST_GROUPS" if month in observed else "MISSING_MONTH")
                report.coverage.append(Coverage(account, month, status))
        except Exception as error:
            # Never persist AWS response bodies, request IDs, raw exception text or credentials.
            detail = str(error) if isinstance(error, DataError) else type(error).__name__
            report.coverage.extend(Coverage(account, month, "QUERY_FAILED", detail) for month in months)
    report.rows.sort(key=lambda row: (row.account, row.month, row.service, row.unit))
    return report


def aggregate(rows: list[CostRow], fields: tuple[str, ...]) -> list[tuple]:
    totals = defaultdict(Decimal)
    estimates = defaultdict(bool)
    for row in rows:
        key = tuple(getattr(row, field) for field in fields) + (row.unit,)
        totals[key] += row.amount
        estimates[key] |= row.estimated
    return [(*key, totals[key], estimates[key]) for key in sorted(totals)]


def write_workbook(report: Report, path: Path, chart: bool = False) -> None:
    """Portable Excel exporter, retaining the source project's openpyxl dependency."""
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)
    account_status = {a: "COMPLETE" if all(c.status == "COMPLETE" for c in report.coverage if c.account == a)
                      else "INCOMPLETE" for a in report.accounts}
    summary = [(a, u, n, estimated, account_status[a])
               for a, u, n, estimated in aggregate(report.rows, ("account",))]
    represented = {row[0] for row in summary}
    summary.extend((a, "n.a.", None, None, "INCOMPLETE") for a in report.accounts if a not in represented)
    monthly = aggregate(report.rows, ("account", "month"))
    services = aggregate(report.rows, ("account", "service"))
    datasets = [
        ("Summary", ["Linked account", "Unit", "Observed cost", "Estimated", "Coverage"], summary),
        ("Monthly Trend", ["Linked account", "Month", "Unit", "Observed cost", "Estimated"], monthly),
        ("Service Breakdown", ["Linked account", "Service", "Unit", "Observed cost", "Estimated"], services),
        ("Coverage", ["Linked account", "Month", "Status", "Issue category"],
         [(c.account, c.month, c.status, c.detail) for c in report.coverage]),
    ]
    for title, headers, records in datasets:
        sheet = book.create_sheet(title)
        sheet.sheet_view.showGridLines = False
        sheet.cell(2, 1, title).font = Font(name="Arial", size=14, bold=True)
        state = "Complete retrieval" if report.complete else "INCOMPLETE: observed amounts are partial"
        label = "SYNTHETIC DEMO. " if report.synthetic else ""
        sheet.cell(3, 1, f"{label}{report.start} to {report.end} (end exclusive). {state}.")
        sheet.cell(4, 1, "UnblendedCost by selected SERVICE names. See Coverage for missing data. Static snapshot.")
        sheet.append([])
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(6, col, header)
            cell.fill = PatternFill("solid", fgColor="232F3E")
            cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for index, record in enumerate(records, 7):
            for col, value in enumerate(record, 1):
                if headers[col - 1] == "Month":
                    value = date.fromisoformat(value)
                elif isinstance(value, bool):
                    value = "Yes" if value else "No"
                cell = sheet.cell(index, col, value)
                if isinstance(value, str):
                    # Treat service names as literal text, even if they start with '='.
                    cell.data_type = "s"
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="center", horizontal="left" if isinstance(value, str) else "right")
                if isinstance(value, Decimal):
                    cell.number_format = '#,##0.00;[Red](#,##0.00);0.00'
                elif isinstance(value, date):
                    cell.number_format = "yyyy-mm"
                elif col == 1:
                    cell.number_format = "@"
        for col, header in enumerate(headers, 1):
            width = 46 if header == "Service" else 28 if header in ("Status", "Issue category", "Coverage") else 20
            sheet.column_dimensions[get_column_letter(col)].width = width
        if len(records) > 12:
            sheet.freeze_panes = "B7"
        sheet.auto_filter.ref = f"A6:{get_column_letter(len(headers))}{max(6, 6 + len(records))}"
    if chart:
        units = {row.unit for row in report.rows}
        if not report.complete or len(units) != 1:
            raise ValueError("chart requires complete coverage and exactly one cost unit")
        sheet = book["Monthly Trend"]
        chart_rows = aggregate(report.rows, ("month",))
        for col, heading in ((8, "Month"), (9, "Observed cost")):
            sheet.cell(6, col, heading)
            sheet.column_dimensions[get_column_letter(col)].width = 20
        for index, (month, unit, amount, _) in enumerate(chart_rows, 7):
            sheet.cell(index, 8, month[:7])
            sheet.cell(index, 9, amount).number_format = "#,##0.00"
        plot = LineChart()
        plot.title = "Selected service costs" + (" (synthetic)" if report.synthetic else "")
        plot.y_axis.title = next(iter(units))
        plot.x_axis.title = "Month"
        plot.legend = None
        plot.add_data(Reference(sheet, min_col=9, min_row=6, max_row=6 + len(chart_rows)), titles_from_data=True)
        plot.set_categories(Reference(sheet, min_col=8, min_row=7, max_row=6 + len(chart_rows)))
        plot.height, plot.width = 10, 20
        sheet.add_chart(plot, "H11")
    with path.open("xb") as output:
        book.save(output)


def write_json(report: Report, path: Path) -> None:
    payload = asdict(report)
    payload.update({"complete": report.complete, "metric": "UnblendedCost",
                    "generated_at_utc": datetime.now(timezone.utc).isoformat()})
    with path.open("x", encoding="utf-8") as output:
        json.dump(payload, output, indent=2, default=str)
        output.write("\n")


class FixtureClient:
    """Replay synthetic CE-shaped pages. Never imports boto3 or contacts AWS."""

    def __init__(self, fixture: dict):
        self.fixture = fixture

    def _response(self, operation: str, request: dict) -> dict:
        filters = request["Filter"]
        dimension = filters["Dimensions"] if "Dimensions" in filters else filters["And"][0]["Dimensions"]
        account = dimension["Values"][0]
        token = request.get("NextPageToken", "0")
        return self.fixture["accounts"][account][operation][int(token)]

    def get_dimension_values(self, **request):
        return self._response("dimensions", request)

    def get_cost_and_usage(self, **request):
        return self._response("costs", request)


def get_ce_client(profile: str, region: str):
    import boto3
    from botocore.config import Config
    session = boto3.Session(profile_name=profile, region_name=region)
    return session.client("ce", config=Config(connect_timeout=10, read_timeout=60,
                          retries={"mode": "standard", "max_attempts": 3}))


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--accounts", type=Path, help="CSV/XLSX with a linked_account_id column")
    source.add_argument("--demo", action="store_true", help="synthetic fixture only, no AWS")
    parser.add_argument("--profile", help="required for live reads, an AWS CLI named profile")
    parser.add_argument("--region", default="us-east-1")
    parser.add_argument("--keywords", default="Bedrock,Claude,Anthropic")
    parser.add_argument("--months", type=int, default=3)
    parser.add_argument("--end-month", help="last completed month to include, YYYY-MM")
    parser.add_argument("--output", type=Path, default=Path("reports/bedrock-cost"), help="output path without extension")
    parser.add_argument("--chart", action="store_true", help="add an Excel monthly line chart when data is complete")
    args = parser.parse_args(argv)
    try:
        if args.demo:
            if args.profile or args.end_month or args.months != 3:
                raise ValueError("demo uses its fixed two-month fixture; omit profile, end-month and months")
            fixture = json.loads((Path(__file__).parent / "fixtures/synthetic_ce.json").read_text())
            accounts, start, end = list(fixture["accounts"]), fixture["start"], fixture["end"]
            keywords = fixture["keywords"]
            ce = FixtureClient(fixture)
        else:
            if not args.profile or not args.profile.strip():
                raise ValueError("live reads require an explicit --profile")
            start, end = reporting_period(args.months, args.end_month)
            accounts = load_accounts(args.accounts)
            keywords = [word.strip() for word in args.keywords.split(",")]
            if not all(keywords):
                raise ValueError("keywords must not contain empty entries")
            ce = None
        outputs = [Path(str(args.output) + extension) for extension in (".xlsx", ".json")]
        if any(path.exists() for path in outputs):
            raise ValueError("output already exists; choose a new --output path")
        args.output.parent.mkdir(parents=True, exist_ok=True)
        if ce is None:
            ce = get_ce_client(args.profile, args.region)
        report = collect_report(ce, accounts, keywords, start, end, synthetic=args.demo)
        # Always retain missing-data evidence, even when a chart cannot be produced.
        use_chart = args.chart and report.complete and len({row.unit for row in report.rows}) == 1
        write_workbook(report, outputs[0], chart=use_chart)
        write_json(report, outputs[1])
        print(f"Saved report: {len(report.rows)} service-month rows, {len(report.coverage)} coverage checks.")
        if args.chart and not use_chart:
            print("Chart omitted: incomplete coverage or multiple/no cost units.")
        print("Complete retrieval." if report.complete else "INCOMPLETE: review Coverage before using observed totals.")
        return 0 if report.complete else 2
    except (ValueError, OSError) as error:
        # Validation text is controlled, but file errors can include local paths.
        print(f"Report not created: {error if isinstance(error, ValueError) else type(error).__name__}", file=sys.stderr)
        return 1
    except Exception as error:
        print(f"Report failed: {type(error).__name__}. No raw service error was logged.", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
