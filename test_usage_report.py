"""Offline regression tests. No credentials, AWS calls or persistent output files."""

import copy
import io
import json
import socket
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook

import usage_report as report


FIXTURE = Path(__file__).resolve().parents[1] / "fixtures/synthetic_ce.json"


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.network = patch.object(socket.socket, "connect", side_effect=AssertionError("network disabled"))
        self.network.start()
        self.addCleanup(self.network.stop)
        self.fixture = json.loads(FIXTURE.read_text())
        self.accounts = list(self.fixture["accounts"])

    def collect(self, ce=None):
        return report.collect_report(ce or report.FixtureClient(self.fixture), self.accounts,
                                     self.fixture["keywords"], self.fixture["start"],
                                     self.fixture["end"], synthetic=True)

    def test_calendar_month_boundaries(self):
        self.assertEqual(report.reporting_period(3, today=date(2026, 3, 31)),
                         ("2025-12-01", "2026-03-01"))
        self.assertEqual(report.reporting_period(1, "2024-02", date(2026, 3, 31)),
                         ("2024-02-01", "2024-03-01"))

    def test_rejects_invalid_months_and_future_partial_month(self):
        for count in (0, -1, 14):
            with self.assertRaises(ValueError):
                report.reporting_period(count)
        for value in ("2026-3", "2026-13", "2026-03", "2027-01"):
            with self.assertRaises(ValueError):
                report.reporting_period(1, value, date(2026, 3, 31))

    def test_strict_account_validation(self):
        self.assertEqual(report.validate_accounts(["000000000001"]), ["000000000001"])
        for values in ([], [111111111111], ["123"], ["١" * 12], ["1" * 12] * 2):
            with self.assertRaises(ValueError):
                report.validate_accounts(values)

    def test_csv_loader_preserves_identifiers_and_rejects_extra_columns(self):
        with tempfile.TemporaryDirectory() as temp:
            file = Path(temp) / "accounts.csv"
            file.write_text("linked_account_id\n000000000001\n")
            self.assertEqual(report.load_accounts(file), ["000000000001"])
            file.write_text("payer_account_id\n111111111111\n")
            with self.assertRaises(ValueError):
                report.load_accounts(file)
            file.write_text("linked_account_id\n111111111111,extra\n")
            with self.assertRaises(ValueError):
                report.load_accounts(file)

    def test_discovery_paginated_and_scoped_per_account(self):
        ce = Mock()
        ce.get_dimension_values.side_effect = self.fixture["accounts"][self.accounts[0]]["dimensions"]
        services = report.discover_services(ce, self.accounts[0], ["BEDROCK", "claude"],
                                           self.fixture["start"], self.fixture["end"])
        self.assertEqual(services, ["Amazon Bedrock", "Synthetic Claude Service"])
        calls = ce.get_dimension_values.call_args_list
        self.assertEqual(len(calls), 2)
        self.assertEqual(calls[0].kwargs["Filter"]["Dimensions"],
                         {"Key": "LINKED_ACCOUNT", "Values": [self.accounts[0]]})
        self.assertEqual(calls[1].kwargs["NextPageToken"], "1")

    def test_pagination_cycle_rejected(self):
        method = Mock(return_value={"NextPageToken": "again"})
        with self.assertRaises(report.DataError):
            list(report.pages(method))
        self.assertEqual(method.call_count, 2)

    def test_populated_views_reconcile_exactly(self):
        result = self.collect()
        self.assertTrue(result.complete)
        self.assertEqual(len(result.rows), 6)
        self.assertEqual(len(result.coverage), 4)
        self.assertEqual(sum(row.amount for row in result.rows), Decimal("36.000"))
        for fields in (("account",), ("account", "month"), ("account", "service")):
            self.assertEqual(sum(row[-2] for row in report.aggregate(result.rows, fields)), Decimal("36.000"))
        self.assertTrue(any(row.estimated for row in result.rows))
        self.assertTrue(any(row.amount == 0 for row in result.rows))
        self.assertTrue(any(row.amount < 0 for row in result.rows))

    def test_cost_queries_always_group_and_filter(self):
        real = report.FixtureClient(self.fixture)
        ce = Mock(wraps=real)
        self.collect(ce)
        self.assertEqual(ce.get_cost_and_usage.call_count, 3)
        for call in ce.get_cost_and_usage.call_args_list:
            self.assertEqual(call.kwargs["GroupBy"], [{"Type": "DIMENSION", "Key": "SERVICE"}])
            self.assertEqual(call.kwargs["Metrics"], ["UnblendedCost"])
            self.assertEqual(call.kwargs["Filter"]["And"][0]["Dimensions"]["Key"], "LINKED_ACCOUNT")

    def test_no_matching_services_is_not_zero_cost(self):
        self.fixture["accounts"][self.accounts[0]]["dimensions"] = [{"DimensionValues": []}]
        result = self.collect()
        self.assertFalse(result.complete)
        self.assertEqual([c.status for c in result.coverage[:2]], ["NO_MATCHING_SERVICES"] * 2)
        self.assertFalse(any(row.account == self.accounts[0] for row in result.rows))

    def test_missing_month_and_empty_groups_are_explicit(self):
        costs = self.fixture["accounts"][self.accounts[1]]["costs"][0]["ResultsByTime"]
        costs[0]["Groups"] = []
        costs.pop()
        result = self.collect()
        self.assertEqual([c.status for c in result.coverage[-2:]], ["NO_COST_GROUPS", "MISSING_MONTH"])
        self.assertFalse(result.complete)

    def test_later_page_failure_discards_partial_account(self):
        real = report.FixtureClient(self.fixture)
        ce = Mock(wraps=real)

        def query(**request):
            if request.get("NextPageToken"):
                raise RuntimeError("private text that must never appear in the report")
            return real.get_cost_and_usage(**request)

        ce.get_cost_and_usage.side_effect = query
        result = self.collect(ce)
        self.assertFalse(any(row.account == self.accounts[0] for row in result.rows))
        self.assertEqual(result.coverage[0].detail, "RuntimeError")
        self.assertNotIn("private text", str(result))

    def test_duplicate_groups_fail_closed(self):
        pages = self.fixture["accounts"][self.accounts[0]]["costs"]
        pages[1]["ResultsByTime"][0]["Groups"] = copy.deepcopy(pages[0]["ResultsByTime"][0]["Groups"])
        result = self.collect()
        self.assertEqual(result.coverage[0].detail, "duplicate_cost_group")

    def test_malformed_metrics_fail_closed(self):
        for bad in ("NaN", "Infinity", "not-a-number"):
            with self.subTest(bad=bad):
                self.fixture["accounts"][self.accounts[1]]["costs"][0]["ResultsByTime"][0]["Groups"][0]["Metrics"]["UnblendedCost"]["Amount"] = bad
                self.assertEqual(self.collect().coverage[-1].status, "QUERY_FAILED")

    def test_currency_units_are_never_summed_together(self):
        self.fixture["accounts"][self.accounts[1]]["costs"][0]["ResultsByTime"][0]["Groups"][0]["Metrics"]["UnblendedCost"]["Unit"] = "EUR"
        result = self.collect()
        totals = report.aggregate(result.rows, ("account",))
        self.assertEqual({row[-3] for row in totals}, {"USD", "EUR"})
        with tempfile.TemporaryDirectory() as temp, self.assertRaises(ValueError):
            report.write_workbook(result, Path(temp) / "report.xlsx", chart=True)

    def test_workbook_contents_chart_and_literal_service_names(self):
        result = self.collect()
        result.rows.append(report.CostRow(self.accounts[0], "2026-01-01", "=1+1", Decimal("0"), "USD", False))
        with tempfile.TemporaryDirectory() as temp:
            file = Path(temp) / "report.xlsx"
            report.write_workbook(result, file, chart=True)
            book = load_workbook(file)
            self.assertEqual(book.sheetnames, ["Summary", "Monthly Trend", "Service Breakdown", "Coverage"])
            self.assertEqual(book["Summary"]["C7"].value, 37.25)
            self.assertEqual(book["Summary"]["A7"].data_type, "s")
            self.assertEqual(book["Service Breakdown"]["B7"].value, "=1+1")
            self.assertEqual(book["Service Breakdown"]["B7"].data_type, "s")
            self.assertEqual(len(book["Monthly Trend"]._charts), 1)
            self.assertIn("SYNTHETIC", book["Summary"]["A3"].value)
            book.close()
            with self.assertRaises(FileExistsError):
                report.write_workbook(result, file)

    def test_demo_cli_is_offline_and_preserves_existing_outputs(self):
        with tempfile.TemporaryDirectory() as temp, patch.object(report, "get_ce_client", side_effect=AssertionError("AWS disabled")):
            stem = Path(temp) / "demo"
            with redirect_stdout(io.StringIO()):
                self.assertEqual(report.main(["--demo", "--chart", "--output", str(stem)]), 0)
            payload = json.loads(stem.with_suffix(".json").read_text())
            self.assertTrue(payload["synthetic"])
            self.assertTrue(payload["complete"])
            self.assertEqual(payload["rows"][0]["amount"], "10.125")
            original = stem.with_suffix(".xlsx").read_bytes()
            with redirect_stderr(io.StringIO()):
                self.assertEqual(report.main(["--demo", "--output", str(stem)]), 1)
            self.assertEqual(stem.with_suffix(".xlsx").read_bytes(), original)

    def test_live_requires_explicit_profile_before_client_creation(self):
        with patch.object(report, "get_ce_client") as client, redirect_stderr(io.StringIO()):
            self.assertEqual(report.main(["--accounts", "unused.csv"]), 1)
            client.assert_not_called()

    def test_incomplete_cli_returns_two_and_writes_coverage(self):
        self.fixture["accounts"][self.accounts[1]]["costs"] = [{"ResultsByTime": []}]
        result = self.collect()
        with tempfile.TemporaryDirectory() as temp, patch.object(report, "collect_report", return_value=result), redirect_stdout(io.StringIO()):
            stem = Path(temp) / "partial"
            self.assertEqual(report.main(["--demo", "--chart", "--output", str(stem)]), 2)
            payload = json.loads(stem.with_suffix(".json").read_text())
            self.assertFalse(payload["complete"])
            self.assertIn("MISSING_MONTH", {row["status"] for row in payload["coverage"]})


if __name__ == "__main__":
    unittest.main()
